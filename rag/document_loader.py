from __future__ import annotations

from io import BytesIO
from pathlib import Path


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_DOCUMENT_BYTES = 12 * 1024 * 1024
MAX_EXTRACTED_CHARS = 2_000_000


def extract_pages(filename: str, content: bytes) -> list[dict]:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Upload PDF, DOCX, or TXT documents only.")
    if not content or len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("Document is empty or exceeds the 12 MB upload limit.")
    if extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("utf-8", errors="replace")
        if len(text) > MAX_EXTRACTED_CHARS: raise ValueError("Extracted text exceeds the 2 million character indexing limit.")
        return [{"page": 1, "text": text}]
    if extension == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install pypdf to process PDF files (pip install pypdf).") from exc
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs cannot be indexed.")
        pages = [{"page": index, "text": page.extract_text() or ""}
                 for index, page in enumerate(reader.pages, start=1)]
        if sum(len(page["text"]) for page in pages) > MAX_EXTRACTED_CHARS:
            raise ValueError("Extracted text exceeds the 2 million character indexing limit.")
        return pages
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Install python-docx to process DOCX files (pip install python-docx).") from exc
    doc = Document(BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    for table in doc.tables:
        text += "\n" + "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    if len(text) > MAX_EXTRACTED_CHARS: raise ValueError("Extracted text exceeds the 2 million character indexing limit.")
    return [{"page": None, "text": text}]


def chunk_pages(pages: list[dict], chunk_words: int = 180, overlap_words: int = 35) -> list[dict]:
    chunks = []
    for page in pages:
        words = (page.get("text") or "").split()
        start = 0
        while start < len(words):
            body = " ".join(words[start:start + chunk_words]).strip()
            if body:
                chunks.append({"page": page.get("page"), "text": body})
            if start + chunk_words >= len(words):
                break
            start += chunk_words - overlap_words
    return chunks
